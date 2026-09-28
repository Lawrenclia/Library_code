"""Native UI integration with synthetic files and a fake backend only."""
import tempfile
import tkinter as tk
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from app import App
from core import HEADERS, Journal, file_hash, read_roster
from operation_log import OperationLog
from settings_panel import DEFAULTS
from tests.test_automation import sample
from tests.test_roster_write import make_roster
from tests.test_wos_import import Bridge


class ImportUITests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.path = self.folder / "list.xlsx"
        make_roster(self.path, flags=(None, None))
        book = load_workbook(self.path)
        for key, value in {"owner": "谭勋策", "title": "Synthetic paper", "doi": "10.1234/test",
                           "matches": 0, "item_ids": "", "reason": ""}.items():
            book.active.cell(2, 2 + list(HEADERS).index(key), value)
        book.save(self.path)
        book.close()
        self.inbox = self.folder / "inbox"
        self.inbox.mkdir()
        self.file = self.inbox / "one.txt"
        self.file.write_bytes(sample())
        self.root = tk.Tk()
        self.root.withdraw()
        with patch("settings_panel.read_preferences", return_value=DEFAULTS):
            self.app = App(self.root, Journal(self.folder / "progress.db"), auto_load=False,
                           operation_log=OperationLog(self.folder / "log.txt"))
        self.app.loaded(read_roster(self.path))
        self.app.owner.set("谭勋策")
        self.app.select_owner()
        self.app.automation_panel.runtime = self.folder / "imports"
        self.panel = self.app.wos_import_panel
        self.panel.inbox.set(str(self.inbox))
        self.bridge = Bridge(self.app.roster.records)
        self.bridge.online = True
        self.bridge.close = lambda: None
        self.app.bridge = self.bridge
        self.runner = patch.object(self.app, "run", side_effect=self.run_sync)
        self.runner.start()

    def tearDown(self):
        self.runner.stop()
        self.app.set_busy(False)
        self.root.update_idletasks()
        self.app.close()
        self.tmp.cleanup()

    def run_sync(self, job, callback, status, **kwargs):
        self.app.set_busy(True)
        try:
            result = job()
        finally:
            self.app.set_busy(False)
        callback(result)

    def test_preview_import_and_restart_do_not_approve_excel(self):
        before = file_hash(self.path)
        self.panel.preview()
        self.assertEqual(self.bridge.calls, [])
        self.assertEqual(self.panel.plan.items[0].status, "ready")
        self.panel.reviewed.set(True)
        with patch("wos_import_panel.messagebox.askyesno", return_value=True):
            self.panel.start()
        self.assertEqual(self.panel.entries["demo-001"]["status"], "pushed")
        self.assertEqual(file_hash(self.path), before)
        self.assertIsNone(self.panel.plan)
        self.assertFalse(self.panel.running)
        self.assertFalse(self.app.busy)
        self.assertFalse(self.panel.reviewed.get())
        self.panel.preview()
        self.assertEqual(self.panel.plan.items[0].status, "pushed")
        self.assertEqual(sum(a == "import_submit" for a, _ in self.bridge.calls), 1)

    def test_no_consent_or_cancel_sends_no_commands(self):
        self.panel.preview()
        with patch("wos_import_panel.messagebox.showwarning") as warning:
            self.panel.start()
            self.assertIn("勾选确认", warning.call_args.args[1])
        self.panel.reviewed.set(True)
        with patch("wos_import_panel.messagebox.askyesno", return_value=False):
            self.panel.start()
        self.assertEqual(self.bridge.calls, [])

    def test_owner_and_limit_changes_invalidate_review(self):
        self.panel.preview()
        self.panel.reviewed.set(True)
        self.panel.limit.set("1")
        self.assertIsNone(self.panel.plan)
        self.assertFalse(self.panel.reviewed.get())
        self.panel.preview()
        self.app.owner.set("另一位")
        self.assertIsNone(self.panel.plan)
        with patch("wos_import_panel.messagebox.showwarning") as warning:
            self.panel.start()
            self.assertIn("不操作其他负责人", warning.call_args.args[1])
        self.assertEqual(self.bridge.calls, [])

    def test_single_local_file_read_needs_no_browser_and_does_not_submit(self):
        self.app.next_record()
        self.app.bridge = None
        with patch("automation_panel.filedialog.askopenfilename", return_value=str(self.file)):
            self.app.automation_panel.load_file()
        state = self.app.automation_panel.store.get(self.app.current)
        self.assertEqual(state["phase"], "exported")
        self.assertEqual(self.bridge.calls, [])

    def test_failed_job_resets_stop_button_and_busy_state(self):
        self.panel.running = True
        self.app.set_busy(True)
        self.assertEqual(str(self.panel.stop_button["state"]), "normal")
        self.app.set_busy(False)
        self.assertFalse(self.panel.running)
        self.assertEqual(str(self.panel.stop_button["state"]), "disabled")
