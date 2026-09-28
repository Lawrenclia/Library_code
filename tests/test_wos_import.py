import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from automation import ImportStore, WOSFlow
from core import SafetyStop
from tests.test_automation import record as base_record, sample, result as sa_result
from wos_import import ImportPlan, build_plan, run_import_plan


def record(**kwargs):
    return base_record(owner="谭勋策", **kwargs)


class Roster:
    def __init__(self, records):
        self.records = records
        self.sha256 = "synthetic-fingerprint"
        self.changed = False

    def assert_unchanged(self):
        if self.changed:
            raise SafetyStop("名单改变")


class Bridge:
    def __init__(self, records):
        self.records = {r.sa_id: r for r in records}
        self.calls = []
        self.batches = {}
        self.fail = None
        self.processed = set()

    def call(self, action, payload, timeout=75):
        self.calls.append((action, payload["sa_id"]))
        if action == self.fail:
            raise SafetyStop("synthetic timeout")
        ident = payload["sa_id"]
        if action == "search":
            r = self.records[ident]
            return sa_result(saLzkId=ident, titleValue=r.title, doiValue=r.doi, wosValue=r.wos,
                             markStatus="已处理" if ident in self.processed else "待处理")
        if action == "import_scan":
            return {"batches": [self.batches[ident]] if ident in self.batches else []}
        if action == "import_upload":
            return {"uploaded": True, "sha256": payload["candidate"]["sha256"]}
        if action == "import_submit":
            self.batches[ident] = dict(id="batch-" + ident, batchNumber="synthetic-" + ident, status=1, modelId="m")
            return {"submitted": True}
        if action == "import_check":
            if ident not in self.batches:
                raise SafetyStop("批次未完成")
            return {"verified": True, "batch": dict(self.batches[ident])}
        if action == "import_push":
            self.batches[ident]["status"] = 2
            return {"submitted": True}
        raise AssertionError("Unexpected command: " + action)


class ImportQueueTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.inbox = self.root / "inbox"
        self.inbox.mkdir()
        self.file = self.inbox / "one.txt"
        self.file.write_bytes(sample())
        self.store = ImportStore(self.root / "imports")
        self.roster = Roster([record()])
        self.bridge = Bridge(self.roster.records)

    def plan(self, **kwargs):
        return build_plan(self.roster, "谭勋策", kwargs.get("limit", 5), self.inbox, self.store)

    def run_plan(self, plan=None, **kwargs):
        return run_import_plan(plan or self.plan(), self.roster, self.bridge, self.store, reviewed=True, **kwargs)

    def actions(self):
        return [a for a, _ in self.bridge.calls]

    def test_plan_is_local_only_and_does_not_stage_import(self):
        plan = self.plan()
        self.assertEqual([item.status for item in plan.items], ["ready"])
        self.assertIsNone(self.store.get(record()))
        self.assertEqual(self.bridge.calls, [])

    def test_owner_scope_done_skipped_and_matches(self):
        self.roster.records += [replace(record(sa_id="other"), owner="另一位"), record(sa_id="done", done=True),
                                record(sa_id="skip", skipped=True), record(sa_id="matched", matches=1)]
        self.assertEqual([i.record.sa_id for i in self.plan().items], ["demo-001"])
        for owner in ("", "另一位"):
            with self.assertRaises(SafetyStop):
                build_plan(self.roster, owner, 5, self.inbox, self.store)
        for limit in (0, 101, True, "5"):
            with self.assertRaises(SafetyStop):
                self.plan(limit=limit)

    def test_requires_explicit_review_and_unchanged_roster(self):
        plan = self.plan()
        with self.assertRaises(SafetyStop):
            run_import_plan(plan, self.roster, self.bridge, self.store)
        self.roster.sha256 = "changed"
        with self.assertRaises(SafetyStop):
            self.run_plan(plan)
        self.assertEqual(self.bridge.calls, [])

    def test_foreign_record_plan_rejected_before_browser(self):
        plan = self.plan()
        item = replace(plan.items[0], record=replace(record(), owner="另一位"))
        with self.assertRaises(SafetyStop):
            self.run_plan(replace(plan, items=(item,)))
        self.assertEqual(self.bridge.calls, [])

    def test_local_file_imports_once_without_wos_or_approval_commands(self):
        result = self.run_plan()
        self.assertEqual(result.outcomes[0]["status"], "pushed")
        for action in ("import_upload", "import_submit", "import_push"):
            self.assertEqual(self.actions().count(action), 1)
        self.assertFalse(any(a.startswith("wos_") for a in self.actions()))
        self.assertEqual(set(self.actions()) - {"search", "import_scan", "import_upload", "import_submit", "import_push", "import_check"}, set())
        self.assertFalse(result.roster.records[0].done)
        second = self.run_plan()
        self.assertEqual(second.outcomes[0]["status"], "pushed")
        self.assertEqual(self.actions().count("import_submit"), 1)

    def test_audit_uses_existing_operation_log_schema(self):
        from operation_log import OperationLog
        log = OperationLog(self.root / "log.txt")
        self.run_plan(audit=log.record)
        events = log.read()
        self.assertIn("import_submit", [event["action"] for event in events])
        self.assertEqual(events[-1]["action"], "WOS 导入队列")
        self.assertEqual(events[-1]["result"], "已执行")

    def test_backend_done_syncs_before_any_import(self):
        self.bridge.processed.add("demo-001")
        updated = Roster([record(done=True)])
        from types import SimpleNamespace
        with patch("wos_import.reconcile_processed", return_value=SimpleNamespace(roster=updated)) as sync:
            result = self.run_plan()
        self.assertEqual(self.actions(), ["search"])
        self.assertEqual(result.outcomes[0]["status"], "synced")
        self.assertIs(result.roster, updated)
        self.assertEqual(sync.call_args.args[2]["markStatus"], "已处理")

    def test_missing_file_still_syncs_already_processed(self):
        self.file.unlink()
        self.bridge.processed.add("demo-001")
        from types import SimpleNamespace
        with patch("wos_import.reconcile_processed", return_value=SimpleNamespace(roster=self.roster)):
            result = self.run_plan()
        self.assertEqual(result.outcomes[0]["status"], "synced")
        self.assertNotIn("import_upload", self.actions())

    def test_no_affiliation_or_weak_identity_never_uploads(self):
        for raw, r in ((sample(C1="Other University"), record()), (sample(), record(doi="")),
                       (sample(DI="10.1234/conflict"), record()), (sample(TI="Different title"), record())):
            with self.subTest(record=r, content=raw[:30]):
                self.file.write_bytes(raw)
                self.roster = Roster([r])
                self.bridge = Bridge(self.roster.records)
                self.assertEqual(self.plan().items[0].status, "deferred")
                result = self.run_plan()
                self.assertEqual(result.outcomes[0]["status"], "deferred")
                self.assertNotIn("import_upload", self.actions())

    def test_changed_file_after_plan_does_not_upload(self):
        plan = self.plan()
        self.file.write_bytes(sample(AF="Another author"))
        self.assertEqual(self.run_plan(plan).outcomes[0]["status"], "deferred")
        self.assertNotIn("import_upload", self.actions())

    def test_distinct_exports_ambiguous_identical_bytes_deduplicated(self):
        other = self.inbox / "copy.txt"
        other.write_bytes(self.file.read_bytes())
        self.assertEqual(self.plan().items[0].status, "ready")
        other.write_bytes(sample(AF="Other author"))
        self.assertEqual(self.plan().items[0].status, "deferred")

    def test_invalid_and_multirecord_files_do_not_replace_valid_file(self):
        (self.inbox / "invalid.txt").write_bytes(b"<html>login</html>")
        raw = sample()
        (self.inbox / "many.txt").write_bytes(raw + raw.split(b"\r\n")[1] + b"\r\n")
        plan = self.plan()
        self.assertEqual(len(plan.file_errors), 2)
        self.assertEqual(plan.items[0].status, "ready")

    def test_duplicate_paper_different_sa_id_never_reimports(self):
        self.roster.records.append(record(sa_id="second-id", row=3))
        self.bridge = Bridge(self.roster.records)
        self.assertEqual([i.status for i in self.plan().items], ["ready", "deferred"])
        self.run_plan()
        next_plan = self.plan()
        self.assertEqual([i.status for i in next_plan.items], ["pushed", "deferred"])
        self.assertEqual(self.actions().count("import_submit"), 1)

    def test_import_timeout_resume_only_reads_existing_submission(self):
        self.bridge.fail = "import_check"
        result = self.run_plan()
        self.assertTrue(result.halted)
        self.assertEqual(self.store.get(record())["phase"], "import_intent")
        self.bridge.fail = None
        self.assertEqual(self.plan().items[0].status, "resume")
        resumed = self.run_plan()
        self.assertEqual(resumed.outcomes[0]["status"], "pushed")
        self.assertEqual(self.actions().count("import_upload"), 1)
        self.assertEqual(self.actions().count("import_submit"), 1)

    def test_uncertain_push_is_not_sent_twice(self):
        self.bridge.fail = "import_push"
        self.assertTrue(self.run_plan().halted)
        self.bridge.fail = None
        self.assertTrue(self.run_plan().halted)
        self.assertEqual(self.actions().count("import_push"), 1)

    def test_uncertain_upload_is_deferred_without_second_upload(self):
        self.bridge.fail = "import_upload"
        self.assertTrue(self.run_plan().halted)
        self.bridge.fail = None
        self.assertEqual(self.plan().items[0].status, "deferred")
        self.run_plan()
        self.assertEqual(self.actions().count("import_upload"), 1)

    def test_disconnection_stops_instead_of_visiting_next_record(self):
        self.roster.records.append(record(sa_id="second-id", row=3))
        self.bridge.fail = "search"
        result = self.run_plan()
        self.assertTrue(result.halted)
        self.assertEqual(len(result.outcomes), 1)
        self.assertEqual(len(self.bridge.calls), 1)

    def test_stop_and_changed_roster_before_run(self):
        result = self.run_plan(stop=lambda: True)
        self.assertTrue(result.cancelled)
        self.assertEqual(self.bridge.calls, [])
        plan = self.plan()
        self.roster.changed = True
        with self.assertRaises(SafetyStop):
            self.run_plan(plan)

    def test_prepare_file_stages_weak_evidence_without_confirming(self):
        flow = WOSFlow(self.bridge, self.store)
        state = flow.prepare_file(record(doi=""), self.file)
        self.assertEqual(state["phase"], "exported")
        self.assertFalse(state["identity_confirmed"])
        self.assertEqual(self.bridge.calls, [])

    def test_existing_archive_cannot_be_silently_replaced(self):
        flow = WOSFlow(self.bridge, self.store)
        flow.prepare_file(record(), self.file)
        self.file.write_bytes(sample(AF="Changed author"))
        with self.assertRaisesRegex(SafetyStop, "已有不同"):
            flow.prepare_file(record(), self.file)

    def test_same_paper_guard_also_protects_existing_single_record_flow(self):
        self.run_plan()
        r = record(sa_id="another-id", row=3)
        flow = WOSFlow(self.bridge, self.store)
        flow.prepare_file(r, self.file)
        with self.assertRaisesRegex(SafetyStop, "同一论文已有导入"):
            flow.proceed(r)
        self.assertEqual(self.actions().count("import_submit"), 1)

    def test_real_synthetic_workbook_sync_keeps_next_import_pending(self):
        from openpyxl import load_workbook
        from core import HEADERS, read_roster
        from tests.test_roster_write import make_roster
        path = self.root / "list.xlsx"
        make_roster(path, flags=(None, None, None))
        book = load_workbook(path)
        sheet = book.active
        columns = {key: 2 + list(HEADERS).index(key) for key in HEADERS}
        for n in (2, 3):
            for key, value in {"owner": "谭勋策", "matches": 0, "item_ids": "", "reason": ""}.items():
                sheet.cell(n, columns[key], value)
        sheet.cell(3, columns["title"], "Synthetic paper")
        sheet.cell(3, columns["doi"], "10.1234/test")
        book.save(path)
        book.close()
        self.roster = read_roster(path)
        self.bridge = Bridge(self.roster.records)
        self.bridge.processed.add("demo-001")
        plan = self.plan()
        self.assertEqual(len(plan.items), 2)
        result = self.run_plan(plan)
        self.assertEqual([o["status"] for o in result.outcomes], ["synced", "pushed"])
        verified = read_roster(path)
        self.assertTrue(verified.records[0].done)
        self.assertFalse(verified.records[1].done)
        self.assertEqual(verified.records[2], self.roster.records[2])
